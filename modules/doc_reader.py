"""Word and Excel document reader"""
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

class DocReader:
    @staticmethod
    def read_word(path: str) -> str:
        try:
            import docx
            doc = docx.Document(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs[:100])
        except ImportError:
            return "python-docx not installed. Run: pip install python-docx"
        except Exception as e:
            return f"Word error: {e}"

    @staticmethod
    def read_excel(path: str, sheet: int = 0) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            ws = wb.worksheets[sheet]
            lines = []
            for row in ws.iter_rows(values_only=True, max_row=50):
                lines.append(" | ".join([str(c) if c is not None else "" for c in row]))
            return "\n".join(lines)
        except ImportError:
            return "openpyxl not installed. Run: pip install openpyxl"
        except Exception as e:
            return f"Excel error: {e}"

    @staticmethod
    def read_csv(path: str) -> str:
        try:
            import csv
            with open(path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                lines = [" | ".join(row) for row in list(reader)[:50]]
                return "\n".join(lines)
        except Exception as e:
            return f"CSV error: {e}"
